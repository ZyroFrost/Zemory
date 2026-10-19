"""move_outputs_to_vm.py — MOVE what is left in data/<task>/{01_raw,02_processing,03_output}/ on the work PC off the PC:
03_output to SharePoint AND the VM, 01_raw / 02_processing to the VM only; then remove the local copy (DW repo docs/plan/22 §7). User 04/10: "dọn là đẩy lên VM nha, ko xóa" · "đẩy VM lẫn đẩy lên SP".

    python scripts\\make_output_links.py            (lists EXTRA files per task)
    python scripts\\move_outputs_to_vm.py           dry run: what would go where
    python scripts\\move_outputs_to_vm.py --apply   do it
    python scripts\\move_outputs_to_vm.py --reference [--apply]
                                                    reference/ only: every file the PC has and the VM lacks or holds
                                                    different; exit 1 when one is missing or differs (no SharePoint login)
    python scripts\\move_outputs_to_vm.py --fix-archive [--apply [--delete-empty]]
                                                    move files an older version put in `<output_url>/_Archive/` up into the
                                                    output_url folder itself (Graph PATCH: item id + version history kept)

Per file, in order — the local copy is removed ONLY when both copies are verified:
  1. SharePoint, folder = `output_url` of tasks/<task>/spec.md, SAME folder the task publishes to (user 04/10: a leftover in
     03_output is an official deliverable that sat in the wrong place, not an archive — "nó vẫn là luồng và file chính thức").
     Same name + same content already there -> nothing to do; missing -> upload; same name + DIFFERENT content -> NOT
     overwritten (it may be the file the task publishes now): reported as CONFLICT, local copy kept, user decides.
     Verified by downloading it back: byte-identical, or for .xlsx identical cell values on every sheet (SharePoint rewrites
     Office files on upload — property promotion — so bytes never match).
     SKIPPED — VM only — for a task that does not publish its own file: `delivers_to: mail`, or an `output_url` in another
     department's folder (it resends the file that department made; pushing there would put a second copy in their folder).
  2. VM, E:\\Automation\\<Repo>\\data\\<Task>\\03_output\\ (the task's working data): same name + same content -> nothing to do;
     same name + DIFFERENT content -> stored as `<stem>__local<ext>` (never overwritten); verified by SHA-256 on the VM.
  3. both verified -> local file removed.
  reference/ (01_raw/reference = templates of an INPUT, 03_output/reference = of the DELIVERABLE: the files a run is rebuilt
  against, written once): COPIED to the same path on the VM, never to SharePoint, verified by SHA-256 — and the local copy
  STAYS. The VM is the backup of the PC; both keep it (user HR 19/10).
Needs: the VM's SSH alias (region value vm_ssh_alias) · SHAREPOINT_* keys in ..\\the region manager\\.env. Only tasks with tasks/<task>/vm_deploy.json.
Shipped by zemory: store part Skill_Task_VmAutomation (its installs). Change the store part, never the repo copy.
"""
import hashlib
import json
import re
import subprocess
import sys
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from region_values import manager_root, value as region_value  # noqa: E402
HOST = region_value(str(Path.cwd()), "vm_ssh_alias", "")
ROOT = Path.cwd()
APPLY = "--apply" in sys.argv
REF_ONLY = "--reference" in sys.argv     # Dept_HR 19/10: report reference/ the VM lacks
REF = {"missing": 0, "differ": 0}
if not (ROOT / "tasks").is_dir():
    sys.exit("run from the repo root (no tasks/ here: %s)" % ROOT)
REPO = ROOT.name
sha = lambda b: hashlib.sha256(b).hexdigest().upper()


# ---------- SharePoint (Graph, app keys of the warehouse repo) ----------
def env_file(p):
    out = {}
    for line in p.read_text(encoding="utf-8").splitlines():
        if "=" in line and not line.lstrip().startswith("#"):
            k, v = line.split("=", 1)
            out[k.strip()] = v.strip().strip('"').strip("'")
    return out


if not HOST or not manager_root(str(ROOT)):
    sys.exit("region not declared: no VM alias or region manager (region_values.py)")
G = "https://graph.microsoft.com/v1.0"
E = {} if REF_ONLY else env_file(Path(manager_root(str(ROOT))) / ".env")
TOK = "" if REF_ONLY else json.loads(urllib.request.urlopen(urllib.request.Request(
    "https://login.microsoftonline.com/%s/oauth2/v2.0/token" % E["SHAREPOINT_TENANT_ID"],
    data=urllib.parse.urlencode({"client_id": E["SHAREPOINT_CLIENT_ID"], "client_secret": E["SHAREPOINT_CLIENT_SECRET"],
                                 "scope": "https://graph.microsoft.com/.default", "grant_type": "client_credentials"}).encode()),
    timeout=60).read())["access_token"]


def g(method, url, data=None, raw=False, ctype=None):
    h = {"Authorization": "Bearer " + TOK}
    if ctype:
        h["Content-Type"] = ctype
    with urllib.request.urlopen(urllib.request.Request(url if url.startswith("http") else G + url, data=data, method=method,
                                                       headers=h), timeout=300) as r:
        b = r.read()
        return b if raw else (json.loads(b) if b else {})


_drives = {}


def sp_target(url):
    """output_url -> (drive id, folder path in the library). https://<host>/sites/<Site>/<Library>/<folder...>"""
    m = re.match(r"https://([^/]+)/sites/([^/]+)/([^/]+)/?(.*)$", urllib.parse.unquote(url))
    host, site, lib, folder = m.groups()
    key = (site, lib)
    if key not in _drives:
        sid = g("GET", "/sites/%s:/sites/%s:/?$select=id" % (host, site))["id"]
        _drives[key] = next(d["id"] for d in g("GET", "/sites/%s/drives?$select=id,webUrl" % sid)["value"]
                            if d["webUrl"].rstrip("/").lower().endswith("/" + lib.lower()))
    return _drives[key], folder.strip("/")


def sp_get(drive, path):
    try:
        it = g("GET", "/drives/%s/root:/%s" % (drive, urllib.parse.quote(path)))
        return it, g("GET", "/drives/%s/items/%s/content" % (drive, it["id"]), raw=True)
    except urllib.error.HTTPError as e:
        if e.code == 404:
            return None, None
        raise


def sp_put(drive, path, data):
    return g("PUT", "/drives/%s/root:/%s:/content" % (drive, urllib.parse.quote(path)), data=data, ctype="application/octet-stream")


# ---------- VM (ssh alias, admin PowerShell) ----------
def ssh_ps(cmd):
    r = subprocess.run(["ssh", HOST, cmd], capture_output=True, text=True, encoding="utf-8", errors="replace")
    return "\n".join(l for l in r.stdout.splitlines() if "post-quantum" not in l and "store now" not in l and "upgraded" not in l).strip()


def vm_hash(path):
    q = path.replace("'", "''")
    out = ssh_ps("if (Test-Path -LiteralPath '%s') { (Get-FileHash -LiteralPath '%s' -Algorithm SHA256).Hash } else { 'MISSING' }" % (q, q))
    return out.splitlines()[-1].strip() if out else "ERROR"


def header(spec):
    t = spec.read_text(encoding="utf-8", errors="replace")
    m = re.match(r"^---\s*\n(.*?)\n---", t, re.S)
    return dict((k.strip(), re.sub(r"\s+#.*$", "", v).strip().strip('"').strip("'"))
                for k, v in (l.split(":", 1) for l in (m.group(1).splitlines() if m else []) if ":" in l and not l.startswith(" ")))


def same_content(name, a, b):
    """Byte-identical, or — for Excel — identical cell values on every sheet. SharePoint rewrites Office files on upload
    (document-property promotion), so an uploaded .xlsx never comes back byte-identical (measured 04/10: 10/10 differ)."""
    if a is None or b is None:
        return False
    if sha(a) == sha(b):
        return True
    if not name.lower().endswith((".xlsx", ".xlsm")):
        return False
    try:
        import io
        import openpyxl
        wa = openpyxl.load_workbook(io.BytesIO(a), data_only=True, read_only=True)
        wb = openpyxl.load_workbook(io.BytesIO(b), data_only=True, read_only=True)
        if wa.sheetnames != wb.sheetnames:
            return False
        return all(list(wa[s].iter_rows(values_only=True)) == list(wb[s].iter_rows(values_only=True)) for s in wa.sheetnames)
    except Exception:
        return False


def local_variant(name):
    p = Path(name)
    return p.stem + "__local" + p.suffix


def task_folder(vm_json):
    """(drive, folder, url) of the first output_url in tasks/<task>/spec.md, or None."""
    spec = vm_json.parent / "spec.md"
    raw_urls = (header(spec).get("output_url") or "") if spec.exists() else ""
    url = next((u.strip() for u in raw_urls.split("|") if u.strip().startswith("http")), None)
    return (sp_target(url) + (url,)) if url else None


if "--fix-archive" in sys.argv:
    # one-off repair of what the _Archive version uploaded: every item of <folder>/_Archive goes up into <folder>
    n_move = n_conf = 0
    for vm_json in sorted(ROOT.glob("tasks/*/vm_deploy.json")):
        t = task_folder(vm_json)
        if not t:
            continue
        drive, folder, url = t
        try:
            arch = g("GET", "/drives/%s/root:/%s/_Archive:/children?$select=id,name,file,folder" % (drive, urllib.parse.quote(folder)))
        except urllib.error.HTTPError as e:
            if e.code == 404:
                continue
            raise
        parent_id = g("GET", "/drives/%s/root:/%s?$select=id" % (drive, urllib.parse.quote(folder)))["id"]
        print("== %s: %s/_Archive (%d item)" % (vm_json.parent.name, folder, len(arch["value"])))
        for it in arch["value"]:
            if "folder" in it:
                print("   SKIP  %-42s sub-folder — by hand" % it["name"]); continue
            main, main_data = sp_get(drive, "%s/%s" % (folder, it["name"]))
            if main is not None:
                _, a_data = sp_get(drive, "%s/_Archive/%s" % (folder, it["name"]))
                same = same_content(it["name"], main_data, a_data)
                print("   %s %-42s a file of that name is already in the folder (%s) — left in _Archive, user decides"
                      % ("DUP  " if same else "CONFL", it["name"], "same content" if same else "DIFFERENT content"))
                n_conf += 1; continue
            if APPLY:
                g("PATCH", "/drives/%s/items/%s" % (drive, it["id"]), data=json.dumps({"parentReference": {"id": parent_id}}).encode(),
                  ctype="application/json")
                ok = sp_get(drive, "%s/%s" % (folder, it["name"]))[0] is not None
                print("   %s %-42s -> %s/" % ("MOVED" if ok else "FAIL ", it["name"], folder))
            else:
                print("   plan  %-42s -> %s/" % (it["name"], folder))
            n_move += 1
        if APPLY:
            left = g("GET", "/drives/%s/root:/%s/_Archive:/children?$select=name" % (drive, urllib.parse.quote(folder)))["value"]
            if not left and "--delete-empty" in sys.argv:          # user 04/10: "xóa _Archive, ko tạo archive gì hết"
                aid = g("GET", "/drives/%s/root:/%s/_Archive?$select=id" % (drive, urllib.parse.quote(folder)))["id"]
                g("DELETE", "/drives/%s/items/%s" % (drive, aid))
                gone = sp_get(drive, "%s/_Archive" % folder)[0] is None
                print("   _Archive (empty) %s" % ("DELETED" if gone else "DELETE FAILED"))
            else:
                print("   _Archive now holds %d item%s" % (len(left), " — empty; --delete-empty removes it" if not left else ""))
    print("%s: %d to move · %d left (name taken in the folder)" % ("APPLIED" if APPLY else "dry run", n_move, n_conf))
    sys.exit(0)

def vm_only_reason(task, spec):
    """Why 03_output of a task goes to the VM only (never SharePoint), or None. A task that mails a file, or resends a
    file another department made (output_url in that department's DataReport folder), publishes nothing of its own."""
    h = header(spec) if spec.exists() else {}
    if (h.get("delivers_to") or "").strip().lower() == "mail":
        return "delivers_to: mail"
    own = task.split("_")[0].upper()
    for u in (h.get("output_url") or "").split("|"):
        m = re.search(r"/DataReport/([^/]+)", urllib.parse.unquote(u.strip()))
        if m and m.group(1).upper() != own:
            return "output_url is the folder of %s" % m.group(1)
    return None


def copy_reference(task, sub):
    """reference/ -> the same path on the VM, verified by SHA-256; the local copy STAYS (the VM is the PC's backup)."""
    src = ROOT / "data" / task / sub / "reference"
    files = sorted(f for f in src.rglob("*") if f.is_file()) if src.is_dir() else []
    if not files:
        return 0
    print("== %s/%s/reference (%d file) -> VM %s, local kept" % (task, sub, len(files), HOST))
    bad = 0
    for f in files:
        part = f.relative_to(src).parent
        rdir = "E:\\Automation\\%s\\data\\%s\\%s\\reference%s" % (REPO, task, sub, "" if str(part) == "." else "\\" + str(part).replace("/", "\\"))
        h = sha(f.read_bytes())
        rh = vm_hash(rdir + "\\" + f.name)
        if rh == h:
            note = "VM: already there (identical)"
        elif rh != "MISSING":
            note, bad = "VM: CONFLICT — %s has DIFFERENT content there, not overwritten (user decides)" % f.name, bad + 1
            REF["differ"] += 1
        elif not APPLY:
            note = "VM: MISSING — would copy"
            REF["missing"] += 1
        else:
            ssh_ps("New-Item -ItemType Directory -Force -Path '%s' | Out-Null" % rdir)
            r = subprocess.run(["scp", "-q", str(f), "%s:%s/" % (HOST, rdir.replace("\\", "/"))], capture_output=True, text=True)
            ok = r.returncode == 0 and vm_hash(rdir + "\\" + f.name) == h
            note, bad = ("VM: copied" if ok else "VM: copied (VERIFY FAILED)"), bad + (0 if ok else 1)
        print("   %s %-42s %s" % ("COPY " if APPLY else "plan ", f.relative_to(src).as_posix(), note))
    return bad


moved = kept = 0
if REF_ONLY:
    for vm_json in sorted(ROOT.glob("tasks/*/vm_deploy.json")):
        for sub in ("01_raw", "03_output"):
            copy_reference(vm_json.parent.name, sub)
    print("reference/: %d missing on the VM · %d different%s" % (REF["missing"], REF["differ"],
          "" if APPLY or not REF["missing"] else " — run again with --apply to copy the missing ones"))
    sys.exit(1 if REF["missing"] or REF["differ"] else 0)
# user 04/10 (new standard): a task that runs on the VM keeps no data on the work PC. Results (03_output) go to SharePoint AND
# the VM; working data (01_raw = what was read · 02_processing = scratch) goes to the VM only — it is nobody's deliverable.
FOLDERS = (("01_raw", False), ("02_processing", False), ("03_output", True))     # (folder, also to SharePoint)
for vm_json in sorted(ROOT.glob("tasks/*/vm_deploy.json")):
    task = vm_json.parent.name
    vm_only = vm_only_reason(task, vm_json.parent / "spec.md")
    for sub, to_sp in FOLDERS:
        if to_sp and vm_only:
            to_sp = False
        if sub in ("01_raw", "03_output"):
            kept += copy_reference(task, sub)
        src = ROOT / "data" / task / sub
        files = sorted(f for f in src.rglob("*") if f.is_file() and not f.name.startswith("_SharePoint")) if src.is_dir() else []
        files = [f for f in files if f.parent == src]                  # top level only; sub-folders listed, moved by hand
        nested = sorted({str(f.parent.relative_to(src)) for f in src.rglob("*") if f.is_file() and f.parent != src}) if src.is_dir() else []
        # reference/ is copied above and stays on the PC (user HR 19/10): not a leftover
        nested = [n for n in nested if not (sub in ("01_raw", "03_output") and Path(n).parts[0] == "reference")]
        if nested:
            print("== %s/%s: sub-folder(s) %s — not handled, move by hand" % (task, sub, ", ".join(nested[:5])))
        if not files:
            continue
        drive = folder = None
        if to_sp:
            t = task_folder(vm_json)
            if not t:
                print("== %s/%s: %d file but spec.md has no output_url — fill it first, nothing moved" % (task, sub, len(files)))
                kept += len(files); continue
            drive, folder, url = t
        rdir = "E:\\Automation\\%s\\data\\%s\\%s" % (REPO, task, sub)
        print("== %s/%s (%d file) -> %sVM %s:%s" % (task, sub, len(files), ("SharePoint %s + " % url) if to_sp else "", HOST, rdir))
        if APPLY:
            ssh_ps("New-Item -ItemType Directory -Force -Path '%s' | Out-Null" % rdir)
        for f in files:
            data = f.read_bytes(); h = sha(data)
            # 1) SharePoint (results only)
            if not to_sp:
                sp_note, sp_ok = "SP: - (%s, VM only)" % (vm_only if sub == "03_output" and vm_only else "working data"), True
            else:
                dest = "%s/%s" % (folder, f.name)
                _, main = sp_get(drive, dest)
                if same_content(f.name, main, data):
                    sp_note, sp_ok = "SP: already in the folder (identical)", True
                elif main is not None:
                    sp_note, sp_ok = "SP: CONFLICT — %s in the folder has DIFFERENT content, not overwritten (user decides)" % f.name, False
                elif not APPLY:
                    sp_note, sp_ok = "SP: would upload into the folder", True
                else:
                    sp_put(drive, dest, data)
                    _, back = sp_get(drive, dest)
                    sp_ok = same_content(f.name, back, data)
                    sp_note = "SP: uploaded into the folder%s" % ("" if sp_ok else " (VERIFY FAILED)")
            # 2) VM
            rname = f.name
            rh = vm_hash(rdir + "\\" + rname)
            if rh not in ("MISSING", h):
                rname = local_variant(f.name)
                rh = vm_hash(rdir + "\\" + rname)
            if rh == h:
                vm_note, vm_ok = "VM: already there (identical)%s" % ("" if rname == f.name else " as " + rname), True
            elif rh != "MISSING":
                vm_note, vm_ok = "VM: %s exists with different content — not overwritten" % rname, False
            elif not APPLY:
                vm_note, vm_ok = "VM: would copy%s" % ("" if rname == f.name else " as " + rname), True
            else:
                tmp = f if rname == f.name else None
                if tmp is None:                                        # copy under the __local name
                    tmp = f.with_name(rname); tmp.write_bytes(data)
                r = subprocess.run(["scp", "-q", str(tmp), "%s:%s/" % (HOST, rdir.replace("\\", "/"))], capture_output=True, text=True)
                if tmp != f:
                    tmp.unlink()
                vm_ok = r.returncode == 0 and vm_hash(rdir + "\\" + rname) == h
                vm_note = "VM: copied%s%s" % ("" if rname == f.name else " as " + rname, "" if vm_ok else " (VERIFY FAILED)")
            # 3) local
            if APPLY and sp_ok and vm_ok:
                f.unlink(); moved += 1; state = "MOVED"
            else:
                state = "plan " if not APPLY else "KEPT "; kept += 0 if not APPLY else 1
            print("   %s %-42s %s · %s" % (state, f.name, sp_note, vm_note))
print("%s: moved %d · kept %d" % ("APPLIED" if APPLY else "dry run", moved, kept))
